# printer.py — заполнение DOCX-шаблона значениями формы (python-docx)
import os
import re
import tempfile

from docx import Document

PLACEHOLDER_RE = re.compile(r"\{\{([a-zA-Z_][a-zA-Z0-9_]*)\}\}")


def _replace_in_paragraph(paragraph, values):
    """Замена {{key}} на уровне параграфа: склеиваем все runs в один текст,
    выполняем замену, результат пишем первым run'ом, остальные очищаем.
    Это нужно, потому что Word может разбить плейсхолдер на несколько runs."""
    text = "".join(run.text for run in paragraph.runs)
    if "{{" not in text:
        return 0
    new_text, n = PLACEHOLDER_RE.subn(
        lambda m: str(values.get(m.group(1), m.group(0))), text)
    if n == 0 or new_text == text:
        return 0
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)
    return n


def fill_docx(template_docx_path, values, output_filename=None):
    """Заполняет DOCX-шаблон словарём {key: value}.

    Возвращает путь к временному файлу результата — вызывающая сторона
    отдаёт его через send_file и удаляет после отправки.
    """
    doc = Document(template_docx_path)
    total = 0
    for paragraph in doc.paragraphs:
        total += _replace_in_paragraph(paragraph, values)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    total += _replace_in_paragraph(paragraph, values)

    fd, out_path = tempfile.mkstemp(prefix="form_", suffix=".docx")
    os.close(fd)
    doc.save(out_path)
    return out_path


def build_download_name(display_name):
    """Имя файла для скачивания: Форма_<display_name>.docx (без запрещённых символов)."""
    safe = re.sub(r'[\\/:*?"<>|]', "_", display_name or "без_имени").strip() or "без_имени"
    return f"Форма_{safe}.docx"
