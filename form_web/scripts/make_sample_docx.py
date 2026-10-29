# scripts/make_sample_docx.py — генерация примера шаблона формы
# Запуск: python scripts/make_sample_docx.py  →  sample_template.docx в корне проекта
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))

from docx import Document


def make_sample(path):
    doc = Document()
    doc.add_heading("Анкета", level=1)
    doc.add_paragraph("ФИО:   {{person}}")
    doc.add_paragraph("ДАТА:  {{date1}}")
    doc.add_paragraph("АДРЕС: {{address}}")
    doc.add_paragraph("ДАТА2: {{date2}}")
    doc.add_paragraph("Телефон: {{phone}}")
    doc.save(path)
    return path


if __name__ == "__main__":
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "sample_template.docx")
    out = os.path.normpath(out)
    make_sample(out)
    print(f"Создан файл: {out}")
    print("Загрузите его на странице /templates")
