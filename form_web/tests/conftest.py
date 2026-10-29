# tests/conftest.py — изолированная БД и uploads для каждого теста
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))


@pytest.fixture()
def app_env(tmp_path, monkeypatch):
    """Создаёт чистую БД и директорию uploads в tmp_path."""
    import database as db
    import app as app_module

    db_file = str(tmp_path / "test.db")
    uploads = str(tmp_path / "uploads")
    os.makedirs(uploads, exist_ok=True)

    monkeypatch.setattr(db, "DB_PATH", db_file)
    monkeypatch.setattr(app_module, "UPLOAD_DIR", uploads)

    db.init_db()
    app_module.app.config["TESTING"] = True
    yield app_module.app
    if os.path.exists(db_file):
        os.remove(db_file)


@pytest.fixture()
def client(app_env):
    return app_env.test_client()


@pytest.fixture()
def conn(app_env):
    import database as db
    c = db.get_connection()
    yield c
    c.close()


def make_docx(path, lines):
    """Вспомогательная генерация DOCX-шаблона с плейсхолдерами."""
    from docx import Document
    doc = Document()
    doc.add_heading("Анкета", level=1)
    for line in lines:
        doc.add_paragraph(line)
    doc.save(path)
    return path
